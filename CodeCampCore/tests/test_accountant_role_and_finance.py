from decimal import Decimal
import io
from django.test import TestCase, Client
from django.contrib.auth.models import User
from django.urls import reverse
import openpyxl

from apps.accounts.models import Profile
from apps.courses.models import Course
from apps.scheduling.models import Batch
from apps.payments.models import Payment, Receipt
from apps.accounts.utils import get_dashboard_url_name


class AccountantRoleAndFinanceTestCase(TestCase):
    def setUp(self):
        self.client = Client()

        # Users
        self.accountant_user = User.objects.create_user(
            username='cpa_david',
            email='david@codecamp.com',
            password='password123',
            first_name='David',
            last_name='Accountant'
        )
        self.accountant_profile = self.accountant_user.profile
        self.accountant_profile.role = 'accountant'
        self.accountant_profile.is_approved = True
        self.accountant_profile.is_verified = True
        self.accountant_profile.onboarding_stage = 'finished'
        self.accountant_profile.save()

        self.student_user = User.objects.create_user(
            username='sarah_student',
            email='sarah@student.com',
            password='password123',
            first_name='Sarah',
            last_name='Debtor'
        )
        self.student_profile = self.student_user.profile
        self.student_profile.role = 'student'
        self.student_profile.is_approved = True
        self.student_profile.is_verified = True
        self.student_profile.onboarding_stage = 'finished'
        self.student_profile.phone = '+2348011223344'
        self.student_profile.save()

        # Course & Batch
        self.course = Course.objects.create(
            name="Full-Stack Cloud Engineering",
            fee=Decimal('200000.00'),
            duration_weeks=16
        )
        self.batch = Batch.objects.create(
            name="Alpha Cohort 2026",
            course=self.course,
            mode='hybrid',
            batch_type='weekdays',
            session_period='morning',
            days_pattern='weekday',
            start_date='2026-10-01',
            end_date='2027-02-01'
        )

        # Payment record (Partial Debtor)
        self.payment = Payment.objects.create(
            student=self.student_user,
            course=self.course,
            batch=self.batch,
            amount_due=Decimal('200000.00'),
            discount=Decimal('20000.00'),
            discount_reason='Early Bird Waiver',
            amount_paid=Decimal('80000.00'),
            status='partial',
            is_approved=True
        )

    def test_accountant_role_and_dashboard_dispatch(self):
        """Accountant profile should route to 'accountant_dashboard'."""
        self.assertEqual(self.accountant_profile.role, 'accountant')
        dashboard_url = get_dashboard_url_name(self.accountant_profile)
        self.assertEqual(dashboard_url, 'accountant_dashboard')

    def test_access_control_gating(self):
        """Students should be redirected away from accountant endpoints, but accountant gets 200 OK."""
        # Student attempting to access accountant dashboard
        self.client.login(username='sarah_student', password='password123')
        resp = self.client.get(reverse('accountant_dashboard'))
        self.assertEqual(resp.status_code, 302)  # Redirected

        resp = self.client.get(reverse('debtors_ledger'))
        self.assertEqual(resp.status_code, 302)  # Redirected

        # Accountant accessing same views
        self.client.login(username='cpa_david', password='password123')
        resp = self.client.get(reverse('accountant_dashboard'))
        self.assertEqual(resp.status_code, 200)

        resp = self.client.get(reverse('debtors_ledger'))
        self.assertEqual(resp.status_code, 200)

    def test_accountant_dashboard_metrics(self):
        """Dashboard should properly aggregate revenue, debtors balance, and collections."""
        self.client.login(username='cpa_david', password='password123')
        resp = self.client.get(reverse('accountant_dashboard'))
        self.assertEqual(resp.status_code, 200)

        # Net expected: 200,000 - 20,000 = 180,000
        # Collected: 80,000
        # Outstanding: 100,000
        self.assertEqual(resp.context['total_gross'], Decimal('200000.00'))
        self.assertEqual(resp.context['total_discounts'], Decimal('20000.00'))
        self.assertEqual(resp.context['net_expected'], Decimal('180000.00'))
        self.assertEqual(resp.context['total_collected'], Decimal('80000.00'))
        self.assertEqual(resp.context['total_outstanding'], Decimal('100000.00'))
        self.assertEqual(resp.context['debtors_count'], 1)

    def test_debtors_ledger_filtering_and_search(self):
        """Debtors ledger should filter debtors, partials, and search terms."""
        self.client.login(username='cpa_david', password='password123')

        # Filter by debtors
        resp = self.client.get(reverse('debtors_ledger'), {'status': 'debtors'})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'Sarah')
        self.assertContains(resp, '100000.00')

        # Search by student username
        resp = self.client.get(reverse('debtors_ledger'), {'q': 'sarah'})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'sarah_student')

    def test_record_direct_payment_and_receipt_issuance(self):
        """Accountant should be able to log a payment and issue an official receipt."""
        self.client.login(username='cpa_david', password='password123')

        record_url = reverse('record_payment')
        resp = self.client.post(record_url, {
            'payment_id': self.payment.id,
            'amount': '100000.00',  # Pay off the remaining balance
            'payment_method': 'Bank Transfer',
            'notes': 'Zenith Bank Ref #ZNTH-9948271',
            'issue_receipt': 'on'
        }, follow=True)

        self.assertEqual(resp.status_code, 200)

        # Refresh from database
        self.payment.refresh_from_db()
        self.assertEqual(self.payment.amount_paid, Decimal('180000.00'))
        self.assertEqual(self.payment.remaining_balance(), Decimal('0.00'))
        self.assertEqual(self.payment.status, 'paid')
        self.assertEqual(self.payment.verified_by, self.accountant_user)
        self.assertTrue(self.payment.is_approved)

        # Verify Receipt was created
        receipt = Receipt.objects.filter(payment=self.payment).order_by('-issued_date').first()
        self.assertIsNotNone(receipt)
        self.assertEqual(receipt.amount, Decimal('100000.00'))

    def test_export_reconciliation_excel(self):
        """Accountant should download a valid .xlsx reconciliation workbook."""
        self.client.login(username='cpa_david', password='password123')

        resp = self.client.get(reverse('export_reconciliation_excel'))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('CodeCamp_Reconciliation_', resp['Content-Disposition'])

        # Verify workbook content using openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(resp.content))
        self.assertIn("Reconciliation", wb.sheetnames)
        ws = wb["Reconciliation"]

        # Check title and header values
        self.assertIn("CODECAMP CORE ACADEMY", ws["A1"].value)
        self.assertEqual(ws["A4"].value, "S/N")
        self.assertEqual(ws["D4"].value, "Student Name")
        self.assertEqual(ws["N4"].value, "Amount Paid (₦)")
        self.assertEqual(ws["O4"].value, "Balance Owed (₦)")

    def test_accountant_receipts_repository(self):
        """Accountant should be able to view all issued receipts."""
        Receipt.objects.create(payment=self.payment, amount=Decimal('80000.00'))

        self.client.login(username='cpa_david', password='password123')
        resp = self.client.get(reverse('accountant_receipts'))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, '80000.00')
        self.assertContains(resp, 'Official Receipts Repository')
